#!/usr/bin/env python3
"""
Gera docs/REQUISITOS.md a partir do DERS (.docx).

O .docx e a FONTE DA VERDADE dos requisitos; o REQUISITOS.md e derivado.

A fonte vive no proprio repositorio, em
docs/Desenvolvimento_DERS_BrunoGabrielGuimaraes.docx -- edite ESSE arquivo no
Word, para a alteracao ser versionada junto com o codigo. Depois rode:

    pip install python-docx
    python docs/importar-ders.py
    python docs/importar-ders.py "C:\\outro\\caminho\\DERS.docx"   # fonte alternativa

O que ele faz:
  1. extrai paragrafos, titulos e tabelas do .docx, na ordem do documento;
  2. extrai as imagens para docs/img/ (so as referenciadas);
  3. monta docs/REQUISITOS.md com as secoes que guiam codigo, preservando a
     numeracao original do DERS (2.2, 2.7.1, RF01, RN01...).

Secoes deliberadamente omitidas, que seguem so no .docx: 1 (Introducao),
2.1 (Os interessados), 2.5 (Estimativas e cronograma), 3.2-3.5 e 4 (Avaliacao,
coberta por VALIDACAO_VERIFICACAO.md).
"""
import io
import os
import re
import shutil
import sys
import zipfile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NOME_DOCX = 'Desenvolvimento_DERS_BrunoGabrielGuimaraes.docx'

# O DERS versionado no repo e a fonte. O caminho em Downloads fica como
# alternativa so para nao quebrar quem ainda edita a copia antiga -- mas essa
# copia nao e versionada, entao o script avisa quando cai nela.
DOCX_REPO = os.path.join(RAIZ, 'docs', NOME_DOCX)
DOCX_ALTERNATIVO = os.path.expanduser(os.path.join('~', 'Downloads', NOME_DOCX))

SAIDA = os.path.join(RAIZ, 'docs', 'REQUISITOS.md')
IMG = os.path.join(RAIZ, 'docs', 'img')
TMP_IMG = os.path.join(RAIZ, 'docs', '.img-tmp')

# (titulo inicial, titulo final exclusivo) na numeracao do DERS
BLOCOS = [
    ('2.2', '2.3'), ('2.3', '2.4'), ('2.4', '2.5'), ('2.6', '2.7'),
    ('2.7', '2.8'), ('2.8', '2.9'), ('2.9', '3'), ('3.1', '3.2'),
]

CABECALHO = '''# Requisitos — SIGAuto

> ⚠️ **ARQUIVO GERADO. NÃO EDITE À MÃO.**
>
> A fonte da verdade dos requisitos é o **DERS em `.docx`**, versionado aqui ao
> lado: [`Desenvolvimento_DERS_BrunoGabrielGuimaraes.docx`](Desenvolvimento_DERS_BrunoGabrielGuimaraes.docx).
> Ele é a entrega acadêmica do TFC — UniRV 2026. Este arquivo é derivado dele, e
> existe para que o código, os testes e os agentes de IA possam ler os requisitos
> sem abrir o Word.
>
> **Mudou um requisito?** Edite o `.docx` **que está nesta pasta** (não uma cópia
> em `Downloads`, que ficaria fora do versionamento) e regenere este arquivo:
>
> ```
> pip install python-docx
> python docs/importar-ders.py
> ```
>
> Qualquer edição feita direto aqui é perdida na próxima importação.
>
> **A numeração original do DERS foi preservada** (2.2, 2.7.1, RF01, RN01…) para
> que a referência cruzada com o documento impresso continue valendo.
>
> **Seções omitidas** por não guiarem código, e que seguem só no `.docx`:
> 1 (Introdução), 2.1 (Os interessados), 2.5 (Estimativas de custo, prazo e
> cronograma), 3.2–3.5 (Arquitetura, gerência, linguagens, persistência — descritos
> no [CLAUDE.md](../CLAUDE.md)) e 4 (Avaliação — ver
> [VALIDACAO_VERIFICACAO.md](VALIDACAO_VERIFICACAO.md)).

## O que NÃO vem daqui

Dois documentos do repositório têm fonte própria, porque só o código pode
verificá-los. Quando o DERS discordar deles, **eles ganham**:

| Arquivo | Fonte | Por quê |
| --- | --- | --- |
| [DER.md](DER.md) | `database/*.sql` | só o schema sabe o que o banco tem; a Figura 26 do DERS é imagem e não se mantém sozinha |
| [VALIDACAO_VERIFICACAO.md](VALIDACAO_VERIFICACAO.md) | o código e as suítes | só rodando os testes se sabe quantos são, e só lendo o código se sabe onde cada RF está |

As divergências já apuradas entre o DERS e o código estão listadas em
[DER.md](DER.md), na última seção.

## Sumário

- [2.2 Restrições obrigatórias](#22-restrições-obrigatórias)
- [2.3 Nomeando convenções e definições](#23-nomeando-convenções-e-definições) — glossário
- [2.4 Fatos e suposições relevantes](#24-fatos-e-suposições-relevantes) — **regras de negócio RN01–RN50**
- [2.6 O escopo do trabalho](#26-o-escopo-do-trabalho) — **escopo incluído e excluído**
- [2.7 Requisitos funcionais](#27-requisitos-funcionais) — **RF01–RF15**
- [2.8 Requisitos não funcionais](#28-requisitos-não-funcionais)
- [2.9 Especificação técnica](#29-especificação-técnica)
- [3.1 Alterações em relação à especificação inicial](#31-requisitos-funcionais)

### Os requisitos funcionais, em uma linha cada

| RF | Nome | RN aplicáveis |
| --- | --- | --- |
| RF01 | Manter clientes | RN06, RN07 |
| RF02 | Manter veículos | RN01, RN11, RN12, RN13 |
| RF03 | Abrir ordem de serviço | RN02, RN09, RN14, RN15, RN16 |
| RF04 | Atualizar ordem de serviço | RN17, RN18, RN22 |
| RF05 | Encerrar ordem de serviço | RN18, RN19, RN20, RN21, RN22 |
| RF06 | Consultar histórico | RN23, RN24 |
| RF07 | Emitir relatórios | RN25, RN26 |
| RF08 | Autenticar usuários | RN03, RN10, RN27, RN28, RN29, RN30 |
| RF09 | Manter estoque | RN31, RN32, RN34, RN35, RN36 |
| RF10 | Movimentar estoque | RN32, RN33 |
| RF11 | Responder via chatbot | RN37, RN38, RN39, RN44 |
| RF12 | Gerar recomendações da IA | RN08, RN40, RN41, RN42, RN43, RN44 |
| RF13 | Manter serviços | RN45, RN46 |
| RF14 | Manter usuários | RN50 |
| RF15 | Manter funcionários | RN47, RN48, RN49 |

> ⚠️ **RF12 é "Gerar recomendações da IA" e RF13 é "Manter serviços".** A matriz de
> rastreabilidade trazia os dois trocados; corrigido conforme este documento.

---

'''


def iter_blocos(doc):
    """Percorre paragrafos e tabelas na ordem real do documento."""
    for child in doc.element.body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            yield Paragraph(child, doc)
        elif tag == 'tbl':
            yield Table(child, doc)


def nivel_titulo(p):
    nome = (p.style.name or '').lower()
    m = re.search(r'heading\s*(\d+)', nome) or re.search(r't[íi]tulo\s*(\d+)', nome)
    if m:
        return int(m.group(1))
    return 1 if nome in ('title', 'título') else None


def extrair(src):
    doc = Document(src)

    rid_para_arquivo = {
        rid: os.path.basename(rel.target_ref)
        for rid, rel in doc.part.rels.items() if 'image' in rel.reltype
    }

    os.makedirs(TMP_IMG, exist_ok=True)
    with zipfile.ZipFile(src) as z:
        for n in z.namelist():
            if n.startswith('word/media/'):
                with z.open(n) as f, open(os.path.join(TMP_IMG, os.path.basename(n)), 'wb') as g:
                    g.write(f.read())

    EMBED = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'

    def texto(p):
        partes = []
        for run in p.runs:
            if run.text:
                t = run.text
                if run.bold and t.strip() and not t.startswith('**'):
                    t = f'**{t}**'
                partes.append(t)
            for el in run._element.iter():
                if el.tag.endswith('}blip'):
                    rid = el.get(EMBED)
                    if rid in rid_para_arquivo:
                        nome = rid_para_arquivo[rid]
                        partes.append(f'\n\n![{nome}](img/{nome})\n')
        return ''.join(partes)

    def celula(c):
        v = ' '.join(p.text.strip() for p in c.paragraphs if p.text.strip())
        return v.replace('|', r'\|').replace('\n', ' ').strip() or ' '

    linhas = []
    for bloco in iter_blocos(doc):
        if isinstance(bloco, Paragraph):
            t = texto(bloco).strip()
            if not t:
                continue
            n = nivel_titulo(bloco)
            if n:
                linhas.append(f'\n{"#" * min(n + 1, 6)} {t.replace("**", "")}\n')
            else:
                estilo = (bloco.style.name or '').lower()
                linhas.append(f'- {t}' if ('list' in estilo or 'lista' in estilo) else t)
                linhas.append('')
        else:
            rows = bloco.rows
            if not rows:
                continue
            ncols = max(len(r.cells) for r in rows)
            head = [celula(c) for c in rows[0].cells] + [' '] * ncols
            linhas += ['', '| ' + ' | '.join(head[:ncols]) + ' |',
                       '|' + '|'.join([' --- '] * ncols) + '|']
            for r in rows[1:]:
                cs = [celula(c) for c in r.cells] + [' '] * ncols
                linhas.append('| ' + ' | '.join(cs[:ncols]) + ' |')
            linhas.append('')

    return '\n'.join(linhas), len(doc.tables)


def _data(caminho):
    """Data de modificacao legivel, para o aviso de copia desatualizada."""
    from datetime import datetime
    return datetime.fromtimestamp(os.path.getmtime(caminho)).strftime('%d/%m/%Y %H:%M')


def escolher_fonte():
    """Caminho do .docx: o argumento, senao a copia do repo, senao Downloads."""
    if len(sys.argv) > 1:
        src = sys.argv[1]
        if not os.path.exists(src):
            sys.exit(f'ERRO: nao achei o DERS em {src}')
        if os.path.abspath(src) != os.path.abspath(DOCX_REPO):
            print('AVISO: gerando a partir de um .docx FORA do repositorio:')
            print(f'       {os.path.abspath(src)}')
            print(f'       A edicao do Word fica fora do versionamento. Copie o arquivo')
            print(f'       para docs/{NOME_DOCX} e edite-o la.')
        return src
    if os.path.exists(DOCX_REPO):
        # A copia do repo existe sempre, entao sem este aviso quem editasse o
        # .docx em Downloads e rodasse sem argumento geraria SILENCIOSAMENTE a
        # partir da copia antiga do repo, com saida de sucesso. Comparar as datas
        # e o que detecta esse caso — era a garantia que faltava.
        if os.path.exists(DOCX_ALTERNATIVO):
            if os.path.getmtime(DOCX_ALTERNATIVO) > os.path.getmtime(DOCX_REPO):
                print('AVISO: a copia de Downloads esta MAIS NOVA que a do repositorio:')
                print(f'       Downloads : {_data(DOCX_ALTERNATIVO)}')
                print(f'       repo      : {_data(DOCX_REPO)}')
                print('       Se voce editou o .docx em Downloads, esta geracao IGNORA essa')
                print(f'       edicao. Copie o arquivo para docs/{NOME_DOCX} primeiro.')
        return DOCX_REPO
    if os.path.exists(DOCX_ALTERNATIVO):
        print('AVISO: usando a copia de Downloads, que NAO e versionada.')
        print(f'       Mova o .docx para docs/{NOME_DOCX} e edite-o la,')
        print('       senao as suas alteracoes ficam fora do repositorio.')
        return DOCX_ALTERNATIVO
    sys.exit(f'ERRO: nao achei o DERS.\n'
             f'  esperado em: {DOCX_REPO}\n'
             f'  ou passe o caminho: python docs/importar-ders.py "C:\\...\\DERS.docx"')


def main():
    # `try/finally` em volta de tudo: `extrair()` descompacta TODAS as midias do
    # .docx (~9 MB) em docs/.img-tmp, e o `rmtree` ficava depois dos `sys.exit`
    # de `achar()` e da checagem de requisitos. Qualquer falha no meio — secao
    # renumerada, que e o caso que o script avisa — deixava o lixo dentro do
    # repositorio, e nao ha entrada para ele no .gitignore.
    try:
        _executar()
    finally:
        shutil.rmtree(TMP_IMG, ignore_errors=True)


def _executar():
    src = escolher_fonte()

    bruto, n_tabelas = extrair(src)
    linhas = bruto.split('\n')

    def achar(prefixo):
        for i, l in enumerate(linhas):
            if re.match(r'^#{2,4}\s+' + re.escape(prefixo) + r'(\s|$)', l):
                return i
        sys.exit(f'ERRO: secao "{prefixo}" nao encontrada no .docx. '
                 f'A numeracao do documento mudou? Ajuste BLOCOS neste script.')

    partes = []
    for ini, fim in BLOCOS:
        a, b = achar(ini), achar(fim)
        if b <= a:
            sys.exit(f'ERRO: ordem invalida, {ini} vem depois de {fim}')
        partes.append('\n'.join(linhas[a:b]).rstrip())

    texto = re.sub(r'\n{4,}', '\n\n\n', CABECALHO + '\n\n'.join(partes) + '\n')

    os.makedirs(IMG, exist_ok=True)
    refs = sorted(set(re.findall(r'\]\(img/([^)]+)\)', texto)))
    copiadas, faltando = [], []
    for nome in refs:
        o = os.path.join(TMP_IMG, nome)
        if os.path.exists(o):
            shutil.copy2(o, os.path.join(IMG, nome))
            copiadas.append(nome)
        else:
            faltando.append(nome)
    shutil.rmtree(TMP_IMG, ignore_errors=True)

    io.open(SAIDA, 'w', encoding='utf-8', newline='\n').write(texto)

    # A contagem olha SO o corpo extraido do .docx, nunca `texto`. O CABECALHO
    # deste script hardcoda RF01-RF15 e 48 das 50 RN: contando sobre `texto`, um
    # RF que desaparecesse do documento continuaria aparecendo na conta e o aviso
    # nunca disparava — so RN04 e RN05 eram detectaveis. A garantia era vazia.
    corpo = '\n\n'.join(partes)
    rf = sorted(set(re.findall(r'RF\d{2}', corpo)))
    rn = sorted(set(re.findall(r'RN\d{2}', corpo)))
    print(f'fonte   : {src}')
    print(f'gerado  : {SAIDA}')
    print(f'tamanho : {len(texto)} caracteres, {len(texto.splitlines())} linhas')
    print(f'tabelas : {n_tabelas} no .docx, {sum(1 for l in texto.splitlines() if l.startswith("|"))} linhas de tabela na saida')
    print(f'imagens : {len(copiadas)} copiadas para docs/img/' + (f' | FALTANDO: {faltando}' if faltando else ''))
    print(f'RF      : {len(rf)} ({rf[0]}..{rf[-1]})' if rf else 'RF      : NENHUM (?)')
    print(f'RN      : {len(rn)} ({rn[0]}..{rn[-1]})' if rn else 'RN      : NENHUM (?)')

    # Requisito que desaparece do .docx faz o script SAIR COM ERRO, nao apenas
    # imprimir aviso: antes o processo saia 0 e nenhum hook ou CI pegava.
    problemas = []
    faltam_rf = {f'RF{i:02d}' for i in range(1, 16)} - set(rf)
    if faltam_rf:
        problemas.append(f'RF ausentes no .docx: {sorted(faltam_rf)}')
    faltam_rn = {f'RN{i:02d}' for i in range(1, 51)} - set(rn)
    if faltam_rn:
        problemas.append(f'RN ausentes no .docx: {sorted(faltam_rn)}')

    if problemas:
        print()
        for p in problemas:
            print(f'ERRO: {p}')
        sys.exit('Requisito faltando: o REQUISITOS.md FOI gravado, mas confira o '
                 '.docx antes de commitar — pode ter sido apagado por acidente.')


if __name__ == '__main__':
    main()
